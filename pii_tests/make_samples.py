"""Generate small DOCX/PPTX files full of fake PII (text, tables, headers, notes, alt text, metadata, an image)."""
import copy
import os
import sys

import docx
import pptx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from pptx.util import Inches as PInches

OUT = os.path.dirname(os.path.abspath(__file__))
IMG = sys.argv[1] if len(sys.argv) > 1 else os.path.join(OUT, "sample_letter.png")


def make_docx(path):
    d = docx.Document()
    d.core_properties.author = "Tobias Lindqvist"
    d.core_properties.last_modified_by = "Helena Vorstermann"
    d.add_heading("Vendor onboarding note", 1)
    p = d.add_paragraph("Contact Priya ")
    p.add_run("Raghavan").bold = True                      # name split across runs
    p.add_run(" at p.raghavan@acme-test.example or +1 (212) 555-0147.")
    d.add_paragraph("Employee Priya Raghavan (EMP-40318), PAN AQBPR7741L, born 14 July 1991.")
    d.add_paragraph("Her home address is 1180 Ashwood Lane, Apartment 4C, Jersey City, NJ 07302 and SSN 900-55-2043.")
    d.add_paragraph("This sentence contains no personal data and must stay untouched.")
    t = d.add_table(rows=2, cols=4)
    for c, v in zip(t.rows[0].cells, ["Name", "Employee ID", "Email", "Phone"]):
        c.text = v
    for c, v in zip(t.rows[1].cells, ["Rohan Deshmukh", "EMP-41488", "r.deshmukh@acme-test.example", "+91 98452 55011"]):
        c.text = v
    d.sections[0].header.paragraphs[0].text = "Policy queries: Delphine Okonjo, d.okonjo@acme-test.example"
    d.sections[0].footer.paragraphs[0].text = "Confidential - prepared by Tobias Lindqvist"
    # mailto hyperlink
    para = d.add_paragraph("Write to ")
    rid = d.part.relate_to("mailto:r.deshmukh@acme-test.example", RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    tt = OxmlElement("w:t")
    tt.text = "r.deshmukh@acme-test.example"
    r.append(tt)
    h.append(r)
    para._p.append(h)
    d.add_paragraph("Scanned request letter:")
    d.add_picture(IMG, width=Inches(5))
    d.inline_shapes[0]._inline.docPr.set("descr", "Letter signed by Mustafa Abdul")
    d.save(path)


def make_pptx(path):
    prs = pptx.Presentation()
    prs.core_properties.author = "Tobias Lindqvist"
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Quarterly vendor review"
    tb = s.shapes.add_textbox(PInches(0.5), PInches(1.5), PInches(8), PInches(1))
    p = tb.text_frame.paragraphs[0]
    p.add_run().text = "Owner: Priya "
    r = p.add_run()
    r.text = "Raghavan"
    r.font.bold = True
    p.add_run().text = " (EMP-40318), p.raghavan@acme-test.example, +1 (212) 555-0147"
    tbl = s.shapes.add_table(2, 3, PInches(0.5), PInches(2.8), PInches(8), PInches(1)).table
    for c, v in zip(tbl.rows[0].cells, ["Name", "Email", "Phone"]):
        c.text = v
    for c, v in zip(tbl.rows[1].cells, ["Rohan Deshmukh", "r.deshmukh@acme-test.example", "+91 98452 55011"]):
        c.text = v
    s.notes_slide.notes_text_frame.text = "Call Delphine Okonjo (d.okonjo@acme-test.example) before the meeting. SSN 900-55-2043."
    pic = s.shapes.add_picture(IMG, PInches(0.5), PInches(4), width=PInches(5))
    pic._element.nvPicPr.cNvPr.set("descr", "Letter signed by Mustafa Abdul")
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Agenda"
    s2.placeholders[1].text = "Budget review\nNo personal data on this slide"
    prs.save(path)


if __name__ == "__main__":
    make_docx(os.path.join(OUT, "sample.docx"))
    make_pptx(os.path.join(OUT, "sample.pptx"))
    print("samples written to", OUT)
